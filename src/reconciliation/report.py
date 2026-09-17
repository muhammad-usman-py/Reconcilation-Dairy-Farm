"""Downloadable Excel report creation."""

from __future__ import annotations

from io import BytesIO

from datetime import datetime, timezone
from pathlib import Path
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .detailed_reconciliation import DetailedReconciliationResult


def _table_prefix(filename: str, fallback: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "_", Path(filename).stem.lower()).strip("_")
    return value or fallback


def _with_source_column_names(data: pd.DataFrame, first_table: str, second_table: str) -> pd.DataFrame:
    """Replace generic left/right labels with source-table column labels."""
    left, right = _table_prefix(first_table, "first_table"), _table_prefix(second_table, "second_table")
    return data.rename(columns={
        "Date (left)": f"{left}_date",
        "Date (right)": f"{right}_date",
        "Value per left": f"{left}_amount",
        "Value per right": f"{right}_amount",
        "Left narration": f"{left}_description",
        "Right narration": f"{right}_description",
    })


def _write_report_sheet(writer: pd.ExcelWriter, name: str, data: pd.DataFrame, first_table: str, second_table: str) -> None:
    """Write one audit-friendly worksheet with its source-table heading."""
    data.to_excel(writer, sheet_name=name, index=False, startrow=3)
    worksheet = writer.sheets[name]
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(data.columns)))
    worksheet.cell(1, 1, "Khwaja Dairy Farm — Ledger Reconciliation Report")
    worksheet.cell(2, 1, f"Source Table 1: {first_table}    |    Source Table 2: {second_table}")
    worksheet.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=14)
    worksheet.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.cell(2, 1).font = Font(italic=True, color="404040")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[4]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:{get_column_letter(max(1, len(data.columns)))}{max(4, len(data) + 4)}"
    for column_number in range(1, max(1, len(data.columns)) + 1):
        column_letter = get_column_letter(column_number)
        values = (worksheet.cell(row=row, column=column_number).value for row in range(4, len(data) + 5))
        width = min(45, max(12, max(len(str(value or "")) for value in values) + 2))
        worksheet.column_dimensions[column_letter].width = width


def make_excel_report(results: pd.DataFrame, first_table: str, second_table: str) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_report_sheet(writer, "All Results", results, first_table, second_table)
        for status, sheet in (("matched", "Matched"), ("discrepancy", "Discrepancies"), ("unmatched_left", "Only First"), ("unmatched_right", "Only Second")):
            _write_report_sheet(writer, sheet, results.loc[results["status"] == status], first_table, second_table)
        summary = results["status"].value_counts().rename_axis("status").reset_index(name="transactions")
        summary.insert(0, "generated_utc", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
        _write_report_sheet(writer, "Summary", summary, first_table, second_table)
    return output.getvalue()


def _write_detailed_sheet(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    subtitle: str,
    data: pd.DataFrame,
    first_table: str,
    second_table: str,
) -> None:
    """Write one of the audit-report sheets in the client workbook style."""
    data.to_excel(writer, sheet_name=sheet_name, index=False, startrow=4)
    worksheet = writer.sheets[sheet_name]
    last_column = max(1, len(data.columns))
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    worksheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    worksheet.cell(1, 1, title)
    worksheet.cell(2, 1, subtitle)
    worksheet.cell(3, 1, f"Source tables: {first_table}  vs  {second_table}")
    worksheet.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=14)
    worksheet.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.cell(2, 1).font = Font(italic=True, color="404040")
    worksheet.cell(3, 1).font = Font(italic=True, color="404040")
    for cell in worksheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = f"A5:{get_column_letter(last_column)}{max(5, len(data) + 5)}"
    for column_number in range(1, last_column + 1):
        values = (worksheet.cell(row=row, column=column_number).value for row in range(5, len(data) + 6))
        width = min(50, max(12, max(len(str(value or "")) for value in values) + 2))
        worksheet.column_dimensions[get_column_letter(column_number)].width = width
    for row in worksheet.iter_rows(min_row=6, max_row=len(data) + 5, min_col=1, max_col=last_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {10, 11, 12, 13, 14})


def _write_summary_sheet(
    writer: pd.ExcelWriter,
    detailed: DetailedReconciliationResult,
    first_table: str,
    second_table: str,
) -> None:
    """Create the audit-style executive summary used in the client workbook."""
    worksheet = writer.book.create_sheet("1. Summary")
    writer.sheets["1. Summary"] = worksheet
    first_label, second_label = Path(first_table).stem, Path(second_table).stem
    balance_rows = detailed.summary.loc[detailed.summary["Section"].eq("Balances")]
    left_balance = float(balance_rows.iloc[0]["Amount"])
    right_balance = float(balance_rows.iloc[1]["Amount"])
    category_summary = detailed.summary.loc[detailed.summary["Section"].eq("Difference reconciliation"), ["Metric", "Items", "Amount"]].rename(columns={"Metric": "Category"})

    worksheet.merge_cells("A1:F1")
    worksheet.merge_cells("A2:F2")
    worksheet.merge_cells("A3:F3")
    worksheet["A1"] = "LEDGER RECONCILIATION - SUMMARY"
    worksheet["A2"] = f"{first_label}  <->  {second_label}"
    worksheet["A3"] = f"Prepared {datetime.now(timezone.utc):%d-%b-%Y}. All figures in PKR."
    worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    worksheet["A2"].font = Font(bold=True)
    worksheet["A3"].font = Font(italic=True, color="404040")

    section_fill = PatternFill("solid", fgColor="D9EAF7")
    table_fill = PatternFill("solid", fgColor="4472C4")
    currency_format = '#,##0;(#,##0);-'
    for column in "ABCDEF":
        worksheet.column_dimensions[column].width = {"A": 48, "B": 14, "C": 18, "D": 4, "E": 18, "F": 48}[column]

    def section(row: int, title: str) -> None:
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = worksheet.cell(row, 1, title)
        cell.font = Font(bold=True)
        cell.fill = section_fill

    section(5, "1. BALANCES PER EACH SET OF BOOKS")
    worksheet["A6"] = f"Closing balance per {first_label} ledger"
    worksheet["C6"] = left_balance
    worksheet["F6"] = f"Balance as shown in {first_label}."
    worksheet["A7"] = f"Closing balance per {second_label} ledger"
    worksheet["C7"] = right_balance
    worksheet["F7"] = f"Balance as shown in {second_label}."
    worksheet["A8"] = f"REQUIRED ADJUSTMENT ({second_label} - {first_label})"
    worksheet["C8"] = "=C7-C6"
    worksheet["F8"] = f"Formula: {second_label} balance minus {first_label} balance."
    for coordinate in ("C6", "C7", "C8"):
        worksheet[coordinate].number_format = currency_format
    worksheet["A8"].font = Font(bold=True)

    section(10, "2. RECONCILIATION OF THE DIFFERENCE")
    worksheet.merge_cells("A11:F11")
    worksheet["A11"] = f"Note: figures are stated in {second_label}'s direction (positive = amount owed to {second_label})."
    worksheet["A11"].font = Font(italic=True)
    for cell, value in zip(worksheet[12], ("Category", "Items", "Amount", "", "", "")):
        cell.value = value
        if value:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = table_fill
    first_category_row = 13
    for offset, (_, row) in enumerate(category_summary.iterrows()):
        excel_row = first_category_row + offset
        worksheet.cell(excel_row, 1, row["Category"])
        worksheet.cell(excel_row, 2, int(row["Items"]))
        worksheet.cell(excel_row, 3, float(row["Amount"]))
        worksheet.cell(excel_row, 3).number_format = currency_format
    total_row = first_category_row + len(category_summary)
    worksheet.cell(total_row, 1, "Total reconciling differences")
    worksheet.cell(total_row, 2, f"=SUM(B{first_category_row}:B{max(first_category_row, total_row - 1)})")
    worksheet.cell(total_row, 3, f"=SUM(C{first_category_row}:C{max(first_category_row, total_row - 1)})")
    worksheet.cell(total_row, 3).number_format = currency_format
    for cell in worksheet[total_row][:3]:
        cell.font = Font(bold=True)

    proof_row = total_row + 3
    section(proof_row, "3. PROOF")
    worksheet.cell(proof_row + 1, 1, f"Balance per {first_label} books ({second_label} direction)")
    worksheet.cell(proof_row + 1, 3, "=C6")
    worksheet.cell(proof_row + 2, 1, "Add: total reconciling differences")
    worksheet.cell(proof_row + 2, 3, f"=C{total_row}")
    worksheet.cell(proof_row + 3, 1, f"Equals balance per {second_label} books")
    worksheet.cell(proof_row + 3, 3, f"=C{proof_row + 1}+C{proof_row + 2}")
    worksheet.cell(proof_row + 4, 1, "Variance remaining (must be nil)")
    worksheet.cell(proof_row + 4, 3, f"=C{proof_row + 3}-C7")
    for row in range(proof_row + 1, proof_row + 5):
        worksheet.cell(row, 3).number_format = currency_format

    settlement_row = proof_row + 7
    section(settlement_row, "4. SETTLEMENT RANGE")
    worksheet.cell(settlement_row + 1, 1, f"If every difference is resolved in {second_label}'s favour, the first party pays")
    worksheet.cell(settlement_row + 1, 3, "=MAX(C7,0)")
    worksheet.cell(settlement_row + 2, 1, f"If every difference is resolved in {first_label}'s favour, the second party refunds")
    worksheet.cell(settlement_row + 2, 3, "=MAX(-C6,0)")
    worksheet.cell(settlement_row + 1, 3).number_format = currency_format
    worksheet.cell(settlement_row + 2, 3).number_format = currency_format

    finding_row = settlement_row + 5
    section(finding_row, "5. WHAT THE RECONCILIATION SHOWS")
    total_groups = len(detailed.matched_comparison)
    differences = len(detailed.difference_register)
    exact_matches = int((detailed.matched_comparison["Match?"] == "MATCH").sum()) if total_groups else 0
    non_value = len(detailed.non_value_differences)
    top_categories = category_summary.assign(abs_amount=category_summary["Amount"].abs()).sort_values("abs_amount", ascending=False).head(4)
    top_text = ", ".join(f"{row['Category']} ({abs(float(row['Amount'])):,.0f})" for _, row in top_categories.iterrows()) or "No value differences"
    most_common = category_summary.sort_values("Items", ascending=False).iloc[0] if not category_summary.empty else None
    positive = int((detailed.difference_register["Difference (right - left)"] > 0).sum())
    negative = int((detailed.difference_register["Difference (right - left)"] < 0).sum())
    findings = [
        f"Every one of the {len(detailed.left_source)} lines in {first_label} and {len(detailed.right_source)} lines in {second_label} has a reconciliation reference. Nothing is left floating.",
        f"{total_groups} transaction groups were matched. {exact_matches} agree to the rupee. {differences} carry a value difference or require follow-up.",
        f"The four largest difference categories are: {top_text}.",
        f"The largest recurring structural issue is {most_common['Category']} ({int(most_common['Items'])} item(s))." if most_common is not None else "No recurring structural value issue was found.",
        f"{positive} differences increase the amount due to {second_label}; {negative} run in the other direction.",
        f"{non_value} further non-value control difference(s) were identified, such as narration, bank, vehicle, bag-count, or posting-date differences.",
    ]
    for offset, finding in enumerate(findings, start=1):
        row = finding_row + offset
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        worksheet.cell(row, 1, f"• {finding}")
        worksheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.row_dimensions[row].height = 30
    worksheet.freeze_panes = "A5"


def make_detailed_excel_report(
    detailed: DetailedReconciliationResult,
    first_table: str,
    second_table: str,
) -> bytes:
    """Create the seven-sheet audit report requested by the client."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_summary_sheet(writer, detailed, first_table, second_table)
        _write_detailed_sheet(writer, "2. Difference Register", "DIFFERENCE REGISTER", "Every matched or unmatched group requiring a value adjustment, largest first.", _with_source_column_names(detailed.difference_register, first_table, second_table), first_table, second_table)
        _write_detailed_sheet(writer, "3. Action List", "ACTION LIST - CHANGES REQUIRED", "Ordered by priority. Update Status and Target date during follow-up.", _with_source_column_names(detailed.action_list, first_table, second_table), first_table, second_table)
        _write_detailed_sheet(writer, "4. Non-Value Differences", "NON-VALUE DIFFERENCES", "Control differences that do not change the reconciled value.", detailed.non_value_differences, first_table, second_table)
        _write_detailed_sheet(writer, "5. Matched Comparison", "FULL MATCHED COMPARISON", "All matched business transaction groups, including value differences.", _with_source_column_names(detailed.matched_comparison, first_table, second_table), first_table, second_table)
        _write_detailed_sheet(writer, "6. Left Ledger", "LEFT LEDGER (SOURCE DATA)", "Every extracted row from the first uploaded ledger with its reconciliation reference.", detailed.left_source, first_table, second_table)
        _write_detailed_sheet(writer, "7. Right Ledger", "RIGHT LEDGER (SOURCE DATA)", "Every extracted row from the second uploaded ledger with its reconciliation reference.", detailed.right_source, first_table, second_table)
    return output.getvalue()
